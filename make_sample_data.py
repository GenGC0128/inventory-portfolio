#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成合成库存横表演示数据（无任何真实业务信息）。"""
import datetime
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).parent / "data" / "reports"
DATES = [datetime.date(2024, 3, 1) + datetime.timedelta(days=i) for i in range(4)]


def sheet_a(path: Path) -> None:
    """模板 A：首行 datetime 日期，第二行 identify，keyword 'qty' 定右边界，total 收尾。"""
    wb, ws = Workbook(), Workbook().active
    ws = wb.active
    ws.append(["report"] + ["", ""] + list(DATES))
    ws.append(["", "", ""] + ["qty", "incoming", "outgoing", "balance"])
    ws.append(["", "", ""])
    ws.append(["", "NovaSilicon", "dense small"] + [120.0, 50.0, 30.0, 140.0])
    ws.append(["", "HelioMaterials", "loose chunk"] + [80.0, 20.0, 10.0, 90.0])
    ws.append(["", "total", ""])
    wb.save(path)


def sheet_b(path: Path) -> None:
    """模板 B：无关键字，右边界到行尾；供应商在第 0 列。"""
    wb, ws = Workbook(), Workbook().active
    ws = wb.active
    ws.append(["label", "", ""] + list(DATES))
    ws.append(["", "", ""] + ["incoming", "outgoing", "balance", "scrap"])
    ws.append(["AuroraPoly", "", "granular type"] + [200.0, 60.0, 20.0, 5.0])
    ws.append(["NovaSilicon", "", "mono chunk"] + [150.0, 40.0, 15.0, 3.0])
    ws.append(["total", "", ""])
    wb.save(path)


def sheet_c(path: Path) -> None:
    """模板 C：Excel 序列日期（从 2024-03-01 起的 serial 值），空行收尾。"""
    wb, ws = Workbook(), Workbook().active
    ws = wb.active
    base_serial = 45352  # 2024-03-01
    ws.append(["", "", ""])
    ws.append(["", "", ""])
    ws.append(["", "", ""] + [base_serial + i for i in range(3)])
    ws.append(["", "", ""] + ["incoming", "outgoing", "balance"])
    ws.append(["TerraWafers", "", "dense block"] + [90.0, 25.0, 95.0])
    ws.append(["HelioMaterials", "", "recycled feed"] + [60.0, 10.0, 55.0])
    ws.append(["", "", ""])
    wb.save(path)


def sheet_project(path: Path) -> None:
    """项目表：行首日期 + 多数量列。"""
    wb, ws = Workbook(), Workbook().active
    ws = wb.active
    ws.append(["date", "ProjX", "ProjY"])
    for i, d in enumerate(DATES):
        ws.append([d, 12.0 * (i + 1), 8.0 * (i + 1)])
    wb.save(path)


def main() -> None:
    for d in ["base_a", "base_b", "base_c", "projects"]:
        (ROOT / d).mkdir(parents=True, exist_ok=True)
    sheet_a(ROOT / "base_a" / "report_a.xlsx")
    sheet_b(ROOT / "base_b" / "report_b.xlsx")
    sheet_c(ROOT / "base_c" / "report_c.xlsx")
    sheet_project(ROOT / "projects" / "project_track.xlsx")
    (ROOT / "base_a" / "corrupted.xlsx").write_bytes(b"not a real xlsx")
    print(f"sample data generated under {ROOT}")


if __name__ == "__main__":
    main()
